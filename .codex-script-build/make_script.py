from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out=Path('presentations/diabetes_risk_screening_15min_script.docx')
out.parent.mkdir(parents=True,exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
NAVY=RGBColor(21,49,58); TEAL=RGBColor(0,139,135); MUTED=RGBColor(89,109,116); CORAL=RGBColor(204,94,81)

def set_font(run,name='Arial',size=11,bold=False,color=NAVY,italic=False):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name); run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),name)
    run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=color

def pformat(p,before=0,after=6,line=1.12):
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=line

styles=doc.styles
for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=styles[style_name]; st.font.name='Arial'; st._element.rPr.rFonts.set(qn('w:ascii'),'Arial'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
styles['Normal'].font.size=Pt(11); styles['Normal'].font.color.rgb=NAVY
styles['Title'].font.size=Pt(28); styles['Title'].font.bold=True; styles['Title'].font.color.rgb=NAVY
styles['Heading 1'].font.size=Pt(17); styles['Heading 1'].font.bold=True; styles['Heading 1'].font.color.rgb=NAVY
styles['Heading 2'].font.size=Pt(12); styles['Heading 2'].font.bold=True; styles['Heading 2'].font.color.rgb=TEAL

# A restrained running footer helps presenters keep loose pages together.
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('CSE 201  |  Diabetes Risk Screening  |  15-minute presenter guide'); set_font(r,size=8,color=MUTED)

p=doc.add_paragraph(style='Title'); pformat(p,after=3); p.add_run('Diabetes Risk Screening Presentation Script')
p=doc.add_paragraph(); pformat(p,after=8)
r=p.add_run('A timed 15-minute guide for the 10-slide presentation'); set_font(r,size=14,color=TEAL)
p=doc.add_paragraph(); pformat(p,after=10)
r=p.add_run('Presenters: '); set_font(r,bold=True); r=p.add_run('[Member 1]  ·  [Member 2]  ·  [Member 3]'); set_font(r)
r=p.add_run('     Class section: '); set_font(r,bold=True); r=p.add_run('[________]'); set_font(r)
p=doc.add_paragraph(); pformat(p,after=5)
r=p.add_run('Pacing: '); set_font(r,bold=True); r=p.add_run('The timings add to 15 minutes. Speak at a steady pace, pause briefly on the charts, and use the transition lines to hand off between presenters. Replace the name and section placeholders before rehearsal.'); set_font(r,color=MUTED)

slides=[
('Slide 1 — Diabetes risk screening','0:45','[Presenter 1]',
'''Good morning, everyone. We are [Member 1], [Member 2], and [Member 3] from [class section]. Our project is a diabetes risk screening demonstration based on the 2015 Behavioral Risk Factor Surveillance System, or BRFSS, Diabetes Health Indicators dataset.

We studied whether survey answers can help identify patterns associated with a respondent reporting a diabetes diagnosis. We built a logistic regression model, compared it with alternatives, and connected the selected model to a Streamlit app. We will explain what the data says, what the model can and cannot do, and what we would improve next.''',
'''Transition: First, we’ll define exactly what goes into the model and what comes out.'''),
('Slide 2 — The problem and target','1:30','[Presenter 1]',
'''The input is a set of 14 answers about a person’s health and background. These include items such as high blood pressure, general health, body mass index, age, and lifestyle indicators. The model combines the answers into a score. We compare that score with a threshold, and the app returns one of two screening messages.

Our target is called Diabetes_binary. A value of 1 means the survey respondent reported having been told they have diabetes. A value of 0 means no diabetes diagnosis in this target. That zero category also includes people classified as prediabetic and people whose diabetes occurred only during pregnancy. So the model learns to distinguish reported diagnosis from this combined comparison group.

This distinction matters. The model predicts patterns in reported diagnoses. It does not test blood sugar, diagnose diabetes, rule it out, or predict who will develop diabetes in the future.''',
'''Transition: The project began with three versions of the source dataset, and their class labels are not identical.'''),
('Slide 3 — The three source datasets','1:30','[Presenter 1]',
'''The original folder contains three BRFSS dataset variants. The first is a three-category file with 253,680 rows. Its target separates no diabetes, prediabetes, and diabetes: 213,703 in category zero, 4,631 in category one, and 35,346 in category two.

The second file is binary but unbalanced. It also has 253,680 rows: 218,334 in class zero and 35,346 in class one. The third is a balanced binary file with 70,692 rows, split evenly between the two classes.

For M1, our project cleaned and deduplicated the balanced data. The resulting modeling dataset has 70,194 rows, with 56,155 for training and 14,039 held out for testing. The held-out people are the common test set used in the model comparisons.''',
'''Transition: With the dataset defined, our EDA notebook asked four questions about its patterns.'''),
('Slide 4 — EDA research questions','1:15','[Presenter 2]',
'''Yes, the EDA notebook explicitly includes research questions, and they are appropriate to present here. The first asks which features are most related to one another. The second asks which features are most associated with the diabetes target. The third asks how reported-diabetes rates vary across age, general health, BMI, education, and income groups. The fourth asks whether the lifestyle indicators show a clear relationship with reported diabetes.

These questions helped us organize the exploratory analysis. EDA means exploratory data analysis: looking for structure, unusual values, and relationships before modeling. It can suggest what deserves attention, but it cannot show that one factor causes another.''',
'''Transition: The clearest pattern was that general health and diagnosed health conditions had stronger associations than the lifestyle measures.'''),
('Slide 5 — EDA findings','1:45','[Presenter 2]',
'''This chart shows the strongest correlations with the binary target. General health has the largest displayed association at 0.38, followed by high blood pressure at 0.36 and BMI at 0.30. High cholesterol, difficulty walking, and age also show positive associations.

The values are moderate. No single feature has a correlation above 0.4, so no one answer cleanly separates respondents with and without a reported diagnosis. The relationship is spread across several answers, which supports using a model that combines features rather than relying on one question.

Lifestyle indicators were weaker in this analysis. Physical activity had a correlation of about negative 0.13, while fruit and vegetable frequency were close to zero. These are descriptive associations in a self-reported survey. They should not be interpreted as proof that a behavior prevents or causes diabetes. For example, health behavior can change after diagnosis, and the survey questions measure only a limited part of someone’s lifestyle.''',
'''Transition: We then built M1 around a fixed training and test split so the final score would not come from the training data.'''),
('Slide 6 — Modeling and evaluation design','1:30','[Presenter 2]',
'''The final cleaned dataset contains 70,194 balanced rows. We used 56,155 rows for training and kept 14,039 separate for testing. The split is stratified, which means the class proportions are preserved in both partitions.

M1 is logistic regression. It uses 14 selected features, and scaling is fitted within the training pipeline. That prevents information from the test rows from influencing preprocessing. The decision threshold is about 0.396. It was selected from out-of-fold predictions during five-fold cross-validation on training data, using 85 percent recall as our working target.

One important limitation is that the fixed test set was later reused for several experiments. Each candidate’s training decisions were kept separate from the test data, but repeated comparisons mean the test is no longer a fresh independent validation. We will describe those later results as exploratory comparisons.''',
'''Transition: Now we’ll look at what M1’s test results mean, especially the balance between finding cases and generating false alarms.'''),
('Slide 7 — M1 test results','2:00','[Presenter 3]',
'''At its selected threshold, M1 reached 85.0 percent recall on the held-out test set. Recall answers: of the people in the positive class, what share did the model flag? Here it found 5,967 of 7,019 positive respondents and missed 1,052.

This is why recall matters for a screening use case. A missed case is someone whose answers resemble the positive class but does not receive a follow-up flag. A screening process may choose to tolerate more false alarms to avoid missing as many possible cases. But that choice has a cost, so we also report specificity and precision.

Specificity was 60.2 percent. Of the 7,020 people in class zero, 4,227 were correctly not flagged and 2,793 were flagged even though they were in the comparison class. Precision on this balanced test set was 68.1 percent. Accuracy was 72.6 percent and ROC AUC was 0.800.

Because the test is balanced, its precision is not a realistic estimate for a population where diabetes prevalence is lower. At an assumed prevalence of 14 percent, our approximate adjustment gives precision around 25.8 percent and a flagged share around 46.1 percent. That is a population-level approximation, not a calibrated personal probability.''',
'''Transition: We also tested whether a much larger or differently regularized model improved the same tradeoff.'''),
('Slide 8 — Model comparisons','2:00','[Presenter 3]',
'''We compared M1 with several alternatives on the same held-out respondents. M1, trained on 56,155 real rows, reached 85.01 percent recall and 60.21 percent specificity.

The saved M2 model used 200,000 training rows, including synthetic diabetes rows. It reached 84.87 percent recall and 59.94 percent specificity, with an AUC of about 0.800. That is not an improvement over M1. A real-only control with about 128,000 rows reached 84.88 percent recall and 60.40 percent specificity, which is very close to M1. This suggests simply adding rows did not create a clear performance gain in this setup.

A new CART-based synthetic-data candidate reached 85.31 percent recall, but specificity fell to 59.07 percent. It found slightly more positives at the cost of more false alarms, so it did not improve our agreed rule. The L2-regularized logistic regression had 84.998 percent recall and 60.21 percent specificity, essentially matching M1 on specificity while narrowly missing the working recall target.

So our experiments do not show that synthetic data improves the model here. The comparisons are useful evidence, but the reused test set means we should not present them as a fresh independent validation.''',
'''Transition: These results informed which model we kept in the app and how we describe its output.'''),
('Slide 9 — The screening app','1:30','[Presenter 3]',
'''The Streamlit app uses M1 and asks the 14 questions required by that model. It combines the answers, compares the score with the selected threshold, and returns one of two broad outcomes: no elevated screening flag, or a suggestion to consider discussing testing with a health professional.

The app’s wording is intentionally cautious. A green or yellow result is not a diagnosis. The data is self-reported, comes from a U.S. survey conducted in 2015, and may not represent people in the Philippines. It may also miss undiagnosed diabetes because the target records reported diagnoses. Optional AI guidance is consent-based and gives general supportive information; it is not medical advice.''',
'''Transition: We’ll close with the decision we can support today and what evidence we still need.'''),
('Slide 10 — What now','1:15','[Presenter 3]',
'''Our recommendation is to keep M1 for the project demonstration. It is the simpler model trained on real data, and the alternatives we tested did not improve the recall-and-specificity rule.

For the next stage, we should keep showing recall together with specificity and the false-alarm burden. We should also avoid making stronger claims from a test set that has been reused. A fresh set of real respondents, ideally representative of the population where the app might be used, would provide better evidence. Finally, our group and instructor should confirm that 85 percent recall is the right working target and decide how many false alarms are acceptable for a screening tool.

Our main conclusion is that the app can demonstrate a data-science workflow, but it is not a clinical screening service. A screening flag may prompt a conversation; diagnosis requires appropriate testing and a qualified health professional. Thank you. We’re ready for your questions.''',
'''End of presentation. Pause for questions.''')
]

for title,time,presenter,spoken,transition in slides:
    h=doc.add_paragraph(style='Heading 1'); pformat(h,before=12,after=3); h.add_run(f'{title}  |  {time}')
    p=doc.add_paragraph(); pformat(p,after=5)
    r=p.add_run('Presenter: '); set_font(r,size=9,bold=True,color=TEAL)
    r=p.add_run(presenter); set_font(r,size=9,color=MUTED)
    for para in spoken.split('\n\n'):
        p=doc.add_paragraph(); pformat(p,after=7,line=1.12); r=p.add_run(para); set_font(r,size=10.5)
    p=doc.add_paragraph(); pformat(p,before=1,after=7); r=p.add_run(transition); set_font(r,size=9.5,bold=True,italic=True,color=TEAL)

h=doc.add_paragraph(style='Heading 1'); pformat(h,before=14,after=5); h.add_run('Rehearsal notes')
for line in [
    'Replace the three member names and class section on the title slide and this guide.',
    'The slide timings total 15 minutes. If rehearsal runs long, shorten the dataset walkthrough and EDA examples before cutting the interpretation of recall and false alarms.',
    'When presenting percentages, say “about” where the slide rounds a stored metric. The 14% prevalence figures are approximate and are not personal risk probabilities.',
    'Use “reported diabetes diagnosis” when describing the target. Avoid saying the model detects diabetes or predicts future diabetes.',
]:
    p=doc.add_paragraph(style='List Bullet'); pformat(p,after=4); r=p.add_run(line); set_font(r,size=10.5)

doc.core_properties.title='Diabetes Risk Screening Presentation Script'
doc.core_properties.subject='Timed 15-minute guide for the CSE 201 diabetes screening project'
doc.core_properties.author='CSE 201 Data Science Fundamentals Project Group'
doc.save(out)
print(f'Saved {out.resolve()}')
print('Word count:',sum(len((title+spoken+transition).split()) for title,tm,pr,spoken,transition in slides))
